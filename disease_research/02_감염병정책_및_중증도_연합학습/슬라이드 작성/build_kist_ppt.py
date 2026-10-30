# -*- coding: utf-8 -*-
"""KIST 연차보고서 → 발표 PPT (소스 deFAI워크샵 포맷 재현)."""
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

# ── 소스 포맷 색/폰트 ───────────────────────────────────────────
NAVY="16263A"; DARK="1C2A38"; BLUE="275BA2"; TEAL="0E7C86"; RED="C0392B"
BOX="F2F7F8"; BOXHI="DCEDEE"; GREENT="1E7A46"; MUT="5B6472"; LINE="D7DEE6"
FB="학교안심 바른돋움 B"; FR="학교안심 바른돋움 R"
SW, SH = 12.5, 7.5

prs=Presentation(); prs.slide_width=I(SW); prs.slide_height=I(SH)
BLANK=prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def _setfont(run,name):
    rPr=run._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        e=rPr.find(qn(tag))
        if e is None:
            e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set('typeface',name)

def tb(s,x,y,w,h,lines,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color, font) ; 한 줄=한 문단"""
    box=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=box.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(1); p.space_before=Pt(0)
        if isinstance(ln,tuple): ln=[ln]
        for (text,size,bold,color,font) in ln:
            r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=RGBColor.from_string(color); r.font.name=font; _setfont(r,font)
    return box

def rect(s,x,y,w,h,fill,line=None,lw=0.75,rounded=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           I(x),I(y),I(w),I(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line: shp.line.color.rgb=RGBColor.from_string(line); shp.line.width=Pt(lw)
    else: shp.line.fill.background()
    shp.shadow.inherit=False
    try: shp.adjustments[0]=0.06
    except Exception: pass
    return shp

def running_title(s,topic):
    # 상단 러닝타이틀 + 주제 바
    tb(s,0.55,0.18,11.4,0.5,[[("KIST 연합 분석·학습 요소기술 ",17,True,NAVY,FB),
                              ("— 추진 실적 및 계획",17,True,MUT,FR)]])
    rect(s,0.55,0.78,0.12,0.42,TEAL,rounded=False)
    tb(s,0.78,0.74,11.3,0.5,[[(topic,18,True,NAVY,FB)]],anchor=MSO_ANCHOR.MIDDLE)
    ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,I(0.55),I(1.34),I(11.4),Pt(1.4))
    ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor.from_string(LINE); ln.line.fill.background(); ln.shadow.inherit=False

def boxrow(s,x,y,w,h,header,body,fill=BOX,hcolor=NAVY,bcolor=DARK,hsize=13,bsize=12.5):
    rect(s,x,y,w,h,fill)
    tb(s,x+0.18,y+0.1,w-0.36,h-0.2,
       [[(header,hsize,True,hcolor,FB)]]+([[(body,bsize,False,bcolor,FR)]] if body else []),
       anchor=MSO_ANCHOR.MIDDLE)

def tag(s,x,y,text,done=True):
    c= GREENT if done else BLUE
    bg="E6F4EC" if done else "E7EEFB"
    r=rect(s,x,y,0.95,0.32,bg,rounded=True);
    tb(s,x,y+0.02,0.95,0.28,[[(text,10.5,True,c,FB)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def stat(s,x,y,w,big,label,color=TEAL):
    rect(s,x,y,w,1.15,BOXHI)
    tb(s,x,y+0.1,w,0.6,[[(big,26,True,color,FB)]],align=PP_ALIGN.CENTER)
    tb(s,x,y+0.72,w,0.4,[[(label,10.5,False,DARK,FR)]],align=PP_ALIGN.CENTER)

# ════════════════════ S1 타이틀 ════════════════════
s=slide()
rect(s,0,0,SW,SH,"FFFFFF",rounded=False)
rect(s,0,2.55,SW,2.4,"F2F7F8",rounded=False)
rect(s,0,2.55,0.22,2.4,TEAL,rounded=False)
tb(s,1.0,2.85,10.5,1.4,[[("KIST 연합학습 과제",40,True,NAVY,FB)],
                        [("추진 실적 및 계획",40,True,TEAL,FB)]])
tb(s,1.0,4.35,10.5,0.5,[[("데이터 유출 없는 프라이버시 보장 연합 분석·학습 요소기술",17,False,DARK,FR)]])
tb(s,1.0,6.6,10.5,0.5,[[("한국과학기술연구원(KIST) · 2026",13,True,MUT,FR)]])

# ════════════════════ S2 개요 ════════════════════
s=slide(); running_title(s,"연구 추진 영역 한눈에 — KIST 5대 업무")
areas=[
 ("① 고속 통신·전송","공동","gRPC 3계층 전송·가중치 압축(float16 2.0×)","UDT/QUIC 비교, SDN/CDN, 캐싱·프리페칭"),
 ("② 프라이버시 보장 분석","단독","동형암호+공개/비밀키 분리, HE 오버헤드 측정","차등프라이버시·시큐어부스트·비민감화"),
 ("③ 외부공격·내부방어","단독","강건집계 방어·악성탐지(AUC 1.0)·워터마킹","트러스트존·분산합의 프로토콜"),
 ("④ 불균형 데이터 정확도","단독","엔트로피 간접추론(r 0.975)·연합 이득 +0.049","전이엔트로피·통합정보량·실데이터"),
 ("⑤ 학습 알고리즘 고도화","단독","FedProx·FedAdam·클라이언트 선택(통신 −81%)","FAVOR/CUCB/FedSAE·지식전파"),
]
y=1.55; rh=1.06
for name,who,done,plan in areas:
    rect(s,0.55,y,11.4,rh-0.12,BOX)
    tb(s,0.75,y+0.08,3.5,0.9,[[(name,14,True,NAVY,FB)],[("KIST 단독" if who=="단독" else "와이즈넛 공동",10,False,MUT,FR)]],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,4.25,y+0.1,4.55,0.85,[[("✓ 한 일  ",10.5,True,GREENT,FB),(done,11.5,False,DARK,FR)]],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,8.95,y+0.1,2.85,0.85,[[("▷ 할 예정  ",10.5,True,BLUE,FB),(plan,10.5,False,DARK,FR)]],anchor=MSO_ANCHOR.MIDDLE)
    y+=rh

# ════════════════════ S3 방법론·수식 ════════════════════
s=slide(); running_title(s,"방법론 · 핵심 수식")
meth=[
 ("예측 모델 + 가중 손실","위중증 예측 MLP(3→64→32→1) · Loss = mean( wᵢ · BCE(ŷᵢ, yᵢ) ),  wᵢ=pos_weight(양성)"),
 ("연합 집계 4종","FedAvg: w = Σₖ (nₖ/N)·wₖ   |   FedProx: +(μ/2)‖w−w_glob‖²   |   FedCurv: ½ Σ Fᵢ(wᵢ−w_glob,ᵢ)²   |   FedAdam: 서버측 Adam"),
 ("동형암호 + 키 분리","Encₚₖ(a) ⊕ Encₚₖ(b) = Encₚₖ(a+b)  ·  복호화는 비밀키(sk, 중앙)만 — 집계자는 합산만 수행"),
 ("불균형·정보이론 지표","라벨 엔트로피 H = −Σ pₖ log₂ pₖ   ·   상호정보량 I(X;Y)  (공유 가능한 스칼라)"),
 ("강건 집계(내부 방어)","좌표별 중앙값 / 절사평균(상·하위 β 제거)  ·  방어 조건: 절사 수 > 악성 수"),
 ("충분통계량 연합(통계형 DiD)","β̂ = (Σₖ Xₖᵀ Wₖ Xₖ)⁻¹ (Σₖ Xₖᵀ Wₖ yₖ)  — 원본 없이 합산만으로 회귀 (그래서 통신 초경량)"),
]
y=1.55; rh=0.9
for h,b in meth:
    rect(s,0.55,y,11.4,rh-0.12,BOX)
    tb(s,0.78,y+0.08,11.0,rh-0.2,[[(h+"   ",12.5,True,NAVY,FB),(b,12,False,DARK,FR)]],anchor=MSO_ANCHOR.MIDDLE)
    y+=rh

# ════════════════════ S4 Ch1 통신 ════════════════════
s=slide(); running_title(s,"① 플랫폼: 고속 통신·전송 (와이즈넛 공동)")
boxrow(s,0.55,1.55,11.4,0.82,"원리","연합학습은 원본 대신 모델 갱신값/통계만 전송 → 병목은 반복 통신. 프로토콜·압축·캐싱으로 최적화.",fill=BOXHI)
tb(s,0.55,2.55,5.0,0.35,[[("✓ 추진 실적",13,True,GREENT,FB)]])
boxrow(s,0.55,2.95,7.4,0.78,"gRPC 기반 3계층 전송 가동","시군구 → 시도 → 중앙 (대용량 채널)")
boxrow(s,0.55,3.81,7.4,0.78,"가중치 압축 시제품","float16 양자화 2.0× · 델타+gzip 1.18×")
boxrow(s,0.55,4.67,7.4,0.78,"통계형(DiD) 연합 초경량","요약통계 전송 3.3KB — 1-shot")
stat(s,8.25,2.95,3.7,"2.0×","float16 양자화 압축비",TEAL)
stat(s,8.25,4.25,3.7,"3.3KB","통계형 연합 1회 전송량",BLUE)
tb(s,0.55,5.62,11.4,0.35,[[("▷ 추진 계획",13,True,BLUE,FB)]])
boxrow(s,0.55,6.0,11.4,0.78,"UDT/QUIC 프로토콜 비교·최적화 · SDN/CDN 경로 최적화 · 캐싱·프리페칭","",fill="EEF3F8")

# ════════════════════ S5 Ch2 프라이버시 ════════════════════
s=slide(); running_title(s,"② 데이터 유출 없는 프라이버시 보장 분석 (KIST 단독)")
boxrow(s,0.55,1.55,11.4,0.82,"원리","동형암호: Encₚₖ(a)⊕Encₚₖ(b)=Encₚₖ(a+b) — 복호화 없이 합산. 공개키(암호화·합산)/비밀키(복호화·중앙) 분리.",fill=BOXHI)
tb(s,0.55,2.55,5.0,0.35,[[("✓ 추진 실적",13,True,GREENT,FB)]])
boxrow(s,0.55,2.95,7.4,0.78,"동형암호(CKKS) 통합 + 공개/비밀키 분리","중간 노드는 내용을 모른 채 합산만")
boxrow(s,0.55,3.81,7.4,0.78,"키 공유 결함 재현·수정","키 1회 생성·분리로 합산·복호화 정상화")
boxrow(s,0.55,4.67,7.4,0.78,"멤버십 추론 공격 검증","모델로부터 참여여부 누출 없음(≈0.50)")
stat(s,8.25,2.95,3.7,"32×","동형암호 암호문 팽창률",BLUE)
stat(s,8.25,4.25,3.7,"0.16 ms","동형 합산 1회 지연(저렴)",TEAL)
tb(s,0.55,5.62,11.4,0.35,[[("▷ 추진 계획",13,True,BLUE,FB)]])
boxrow(s,0.55,6.0,11.4,0.78,"차등프라이버시 모듈 통합·강도-정확도 곡선 · 데이터 비민감화·레이블 유출 방지·시큐어부스트","",fill="EEF3F8")

# ════════════════════ S6 Ch3 보안 ════════════════════
s=slide(); running_title(s,"③ 외부공격·내부방어 기반 보안 (KIST 단독)")
boxrow(s,0.55,1.42,11.4,0.6,"원리","트러스트존(신뢰구역 차등권한) · 분산합의(무결성 교차검증) · 신경망 워터마킹(모델 소유권)",hsize=12.5,bsize=12)
# 공격 정의 박스
rect(s,0.55,2.1,11.4,0.72,BOXHI)
tb(s,0.78,2.18,11.0,0.58,[[("공격 = 업데이트 증폭(모델 포이즈닝)   ",12.5,True,RED,FB),
   ("악성 노드 4/21개가 라벨을 뒤집어 학습한 뒤, 자기 가중치 변화량을 ×6로 부풀려 전송 → 평균 집계를 오염시키려는 내부 공격.  ‘방어’ = 이런 변조에도 전역 모델 정확도를 정상(무공격) 수준으로 유지하는 것.",11.5,False,DARK,FR)]],anchor=MSO_ANCHOR.MIDDLE)
tb(s,0.55,2.9,7.0,0.3,[[("✓ 집계 방식별 방어 결과 — 전역 AUC (높을수록 정상)",12,True,GREENT,FB)]])
# 막대 차트
cd=CategoryChartData(); cd.categories=["공격없음","평균(FedAvg)","중앙값(방어)","절사평균"]
cd.add_series("AUC",(0.823,0.819,0.823,0.200))
gf=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,I(0.55),I(3.22),I(6.7),I(2.72),cd).chart
gf.has_legend=False; gf.has_title=False
gf.value_axis.minimum_scale=0; gf.value_axis.maximum_scale=0.9
plot=gf.plots[0]; plot.has_data_labels=True; plot.data_labels.number_format='0.00'; plot.data_labels.number_format_is_linked=False
plot.data_labels.font.size=Pt(11); plot.data_labels.font.bold=True
for ax in (gf.category_axis,gf.value_axis): ax.tick_labels.font.size=Pt(10)
stat(s,7.55,3.22,4.4,"AUC 1.000","악성 노드 탐지(업데이트 거리)","C0392B")
# 집계 방식의 뜻
rect(s,7.55,4.52,4.4,1.42,BOX)
tb(s,7.75,4.62,4.05,1.26,[
  [("집계 방식의 뜻",12,True,NAVY,FB)],
  [("• 평균(FedAvg): 표본수 가중 평균",10.5,False,DARK,FR)],
  [("• 중앙값: 좌표별 중앙값 — 극단값 무시(방어 성공)",10.5,False,TEAL,FR)],
  [("• 절사평균: 상·하위 잘라낸 평균 — 절사 수<악성 수면 실패",10.5,False,DARK,FR)],
])
boxrow(s,0.55,6.04,11.4,0.42,"신경망 워터마킹  —  본임무 유지(0.823→0.810) · 트리거 적중 0%→100% 소유권 식별","",hsize=11.5)
boxrow(s,0.55,6.5,11.4,0.46,"▷ 할 예정: 트러스트존 운영체계 · 분산합의 프로토콜 · 워터마킹 강건화","",fill="EEF3F8",hsize=11.5)

# ════════════════════ S7 Ch4 불균형 ════════════════════
s=slide(); running_title(s,"④ 단말·사일로 간 불균형 데이터 정확도 (KIST 단독)")
boxrow(s,0.55,1.5,11.4,0.72,"원리","공유 가능한 스칼라(라벨 엔트로피·상호정보량)만으로 원본 없이 사일로 불균형·난이도를 간접 추론.",fill=BOXHI)
tb(s,0.55,2.42,6.0,0.32,[[("✓ 추진 실적 — 엔트로피 ↔ 단독학습 F1 (21개 구)",12.5,True,GREENT,FB)]])
xy=XyChartData(); sname=xy.add_series("구")
pts=[(0.430,0.3478),(0.575,0.4611),(0.732,0.5930),(0.866,0.6800),(0.943,0.7033),(0.425,0.3381),(0.577,0.4469),
     (0.728,0.5231),(0.839,0.5960),(0.935,0.6692),(0.434,0.2769),(0.560,0.4160),(0.701,0.6053),(0.858,0.6731),
     (0.928,0.6818),(0.420,0.3597),(0.527,0.4339),(0.733,0.5120),(0.843,0.6400),(0.938,0.7419),(0.386,0.3120)]
for x,yv in pts: sname.add_data_point(x,yv)
ch=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER,I(0.55),I(2.78),I(6.7),I(3.35),xy).chart
ch.has_legend=False; ch.has_title=False
ch.category_axis.minimum_scale=0.35; ch.category_axis.maximum_scale=0.98
ch.value_axis.minimum_scale=0.25; ch.value_axis.maximum_scale=0.78
ch.category_axis.tick_labels.font.size=Pt(9); ch.value_axis.tick_labels.font.size=Pt(9)
ch.category_axis.has_title=True; ch.category_axis.axis_title.text_frame.text="라벨 엔트로피 (불균형 낮음 → 높음)"
ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
ser=ch.series[0]; ser.marker.style=8; ser.marker.size=7
ser.format.fill.solid(); ser.format.fill.fore_color.rgb=RGBColor.from_string(BLUE)
tb(s,4.85,3.05,2.4,0.4,[[("상관 r = +0.975",13,True,TEAL,FB)]])
stat(s,7.55,2.78,4.4,"+0.049","희소 구간: 연합 vs 단독학습(평균)",TEAL)
boxrow(s,7.55,4.08,4.4,1.0,"최약체 사일로 ↑ = 연합 참여 효과","혼자 학습 0.682 → 연합 참여 0.818 (+0.136). 협력 자체가 약자를 끌어올림(전용 보정 아님)",hsize=11.5,bsize=10.5)
boxrow(s,7.55,5.18,4.4,0.92,"성과 = 식별 + 끌어올림","① 엔트로피로 취약 사일로 식별  ② 연합으로 약체 향상 (약체 표적 가중집계는 향후)",fill="EEF3F8",hsize=11,bsize=10)
boxrow(s,0.55,6.35,11.4,0.62,"▷ 할 예정: 약체 표적 가중집계 · 전이 엔트로피·통합 정보량(인과 흐름) · 실증분야 데이터","",fill="EEF3F8",hsize=12)

# ════════════════════ S8 Ch5 알고리즘 ════════════════════
s=slide(); running_title(s,"⑤ 연합 학습 알고리즘 고도화 (KIST 단독)")
boxrow(s,0.55,1.5,11.4,0.72,"원리","클라이언트 선택(일부 단말만 참여→통신↓) · FedProx/FedCurv 근접항으로 분포 차이(non-IID) 완화.",fill=BOXHI)
tb(s,0.55,2.42,7.0,0.32,[[("✓ 추진 실적 — 총 통신 횟수 비교 (정확도는 동률)",12.5,True,GREENT,FB)]])
cd=CategoryChartData(); cd.categories=["전체 (21구/라운드)","선택 (4구/라운드)"]
cd.add_series("누적 모델 업데이트",(420,80))
bch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,I(0.55),I(2.85),I(6.7),I(3.05),cd).chart
bch.has_legend=False; bch.has_title=False
bch.value_axis.minimum_scale=0; bch.value_axis.maximum_scale=480
bp=bch.plots[0]; bp.has_data_labels=True; bp.data_labels.font.size=Pt(14); bp.data_labels.font.bold=True
bp.data_labels.number_format='0"회"'; bp.data_labels.number_format_is_linked=False
bch.category_axis.tick_labels.font.size=Pt(11); bch.value_axis.tick_labels.font.size=Pt(9)
pts=bch.series[0].points
pts[0].format.fill.solid(); pts[0].format.fill.fore_color.rgb=RGBColor.from_string("9AA7B4")
pts[1].format.fill.solid(); pts[1].format.fill.fore_color.rgb=RGBColor.from_string(TEAL)
tb(s,3.95,3.2,3.1,0.85,[[("정확도(AUC)",12,True,DARK,FR)],[("둘 다 ≈ 0.823",14,True,TEAL,FB)],[("— 동률 ✓",12,True,TEAL,FB)]],align=PP_ALIGN.CENTER)
stat(s,7.55,2.78,4.4,"−81%","통신 '횟수' 절감 (압축과 별개)",TEAL)
boxrow(s,7.55,4.08,4.4,1.0,"어떻게? = 클라이언트 선택","라운드마다 21구 중 4구만 참여 → 전송 횟수 420→80회 (1−4/21≈81%↓), 정확도 0.823 동일",hsize=11.5,bsize=10.5)
boxrow(s,7.55,5.18,4.4,0.92,"압축과는 다른 축 (곱해짐)","Ch1 압축(byte −50%, float16)과 결합 시 ≈ −90%까지",fill="EEF3F8",hsize=11,bsize=10.5)
boxrow(s,0.55,6.35,11.4,0.62,"▷ 할 예정: FAVOR/CUCB/FedSAE 지능형 선택 · FedCurv 완전형 · 연합 대리학습·지식 전파","",fill="EEF3F8",hsize=12)

# ════════════════════ S9 종합 계획 ════════════════════
s=slide(); running_title(s,"향후 계획 (2차년도 추진 방향)")
plans=[
 ("고속 통신·전송 (공동)","UDT/QUIC 프로토콜 비교·최적화 · SDN/CDN 경로 최적화 · 델타압축 고도화·캐싱·프리페칭"),
 ("프라이버시 보장 분석","차등프라이버시 통합 및 강도-정확도 곡선 · 데이터 비민감화·레이블 유출 방지·시큐어부스트"),
 ("외부공격·내부방어","Assistant–중앙 트러스트존 운영체계 · 분산합의 무결성 프로토콜 · 워터마킹 강건화"),
 ("불균형 데이터 정확도","전이 엔트로피·통합 정보량 기반 인과 흐름 측정 · 균형 검증셋 · 실증분야(실데이터) 적용"),
 ("학습 알고리즘 고도화","FAVOR/CUCB/FedSAE 지능형 클라이언트 선택 · FedCurv 완전형 · 연합 대리학습·지식 전파"),
]
y=1.55; rh=1.06
for h,b in plans:
    rect(s,0.55,y,11.4,rh-0.12,BOX); rect(s,0.55,y,0.1,rh-0.12,BLUE,rounded=False)
    tb(s,0.8,y+0.12,11.0,rh-0.28,[[(h,13.5,True,NAVY,FB)],[(b,12,False,DARK,FR)]],anchor=MSO_ANCHOR.MIDDLE)
    y+=rh

prs.save("KIST_연차보고서_발표.pptx")
print("saved: KIST_연차보고서_발표.pptx |", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
