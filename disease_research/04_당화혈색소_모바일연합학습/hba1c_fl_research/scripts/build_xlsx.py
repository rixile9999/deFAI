from pathlib import Path
RESEARCH_ROOT = Path(__file__).resolve().parents[1]
(RESEARCH_ROOT / "report").mkdir(exist_ok=True)
# -*- coding: utf-8 -*-
import sys; sys.path.insert(0,str(RESEARCH_ROOT / 'scripts'))
import pandas as pd
from candidates import C
from openpyxl.styles import Alignment, Font, PatternFill
cols=[('ID','ID'),('name','공식명칭'),('org','제공기관'),('country','국가'),('version','버전/수집기간'),('type','자료유형'),('cls','분류(①~④)'),('level','확인수준'),('url','공식 소개·다운로드·신청 페이지'),('paper','원논문/공식문서'),('access','접근조건·절차'),('cost','비용'),('redistrib','이용허락·재배포·학습제한'),('n_reported','보고 표본수'),('n_usable','필요변수 모두 있는 실제 사용가능 표본수'),('repeat','반복검사·추적'),('hba1c','HbA1c 변수·단위·측정법'),('glucose','혈당 종류·변수'),('weight','체중 변수·측정시점'),('other','기타 변수(CGM·활동·임상상태)'),('link','사람·방문·시점 연결/구분자'),('missing','결측·분포'),('A','A 현재추정 적합성'),('B','B 미래예측 적합성'),('C','C 개인별 FL·개인화 적합성'),('status','확보상태'),('note','비고·한계')]
df=pd.DataFrame([{k2:c.get(k,'') for k,k2 in cols} for c in C])
V=[
# name, level, dtype, unit, source, missing code, QC rule, note
('participant_pid','필수','string','—','연구 발급 가명ID(앱 설치 시 생성, 키는 기관 분리보관)','불허(NOT NULL)','UUID 형식, 1인 1ID, 재등록 시 연결테이블','실명·전화번호 저장 금지'),
('site_id','필수(cross-silo)/권장','string','—','모집 기관 코드','"UNK"','사전 정의 코드 목록','기관단위 FL·층화'),
('device_id','필수','string','—','앱 설치/센서 일련번호의 해시','"UNK"','기기 교체 시 신규 device_id + 교체일 기록','기기교체 분석'),
('visit_id','필수','string','—','연구 방문 ID','불허','participant_pid+순번 유일',''),
('lab_test_id','필수','string','—','검사 1건 ID','불허','중복 업로드 탐지(동일값·동일일)',''),
('hba1c_value_raw','필수','decimal','원본 단위 그대로','검사결과지/LIS','공란=NA, "측정불가"는 별도 코드','원본 문자열 보존; 숫자 변환 실패 플래그','원본값 덮어쓰기 금지'),
('hba1c_unit_raw','필수','enum','%(NGSP) | mmol/mol(IFCC)','결과지','불허','허용 단위 외 값 거부',''),
('hba1c_pct_std','필수(파생)','decimal(1)','% NGSP','변환: NGSP%=0.09148×IFCC+2.152','NA','범위 3.0–20.0 외 검토 플래그','마스터방정식(NGSP/IFCC)'),
('hba1c_mmolmol_std','권장(파생)','decimal(0)','mmol/mol','IFCC=10.93×NGSP%−23.50','NA','',''),
('hba1c_date','필수','date','YYYY-MM-DD','결과지 채혈일(보고일 아님)','불허(없으면 행 제외)','채혈일≤보고일≤업로드일','시간연결의 기준점'),
('hba1c_method','권장','enum','HPLC|면역법|효소법|POC|DBS|미상','결과지/기관 문의','"미상"','NGSP 인증 여부 별도 필드','혈색소변이 간섭은 방법별 상이'),
('hba1c_source','필수','enum','검사실(정맥)|현장검사POC|가정채혈키트|자가보고','수집 경로','불허','자가보고는 검증 전 학습 제외','출처확인'),
('hba1c_verification','필수','enum','결과지 사진 대조|기관 LIS 연동|미검증','연구 코디네이터','"미검증"','이중입력 불일치 시 재확인',''),
('is_estimated_a1c','필수','bool','—','eAG/GMI 등 추정값 여부','불허','TRUE면 정답(label) 사용 금지','실측과 혼합 금지'),
('glucose_value_raw','필수','decimal','원본 단위','혈당측정기/CGM/검사실','NA','',''),
('glucose_unit_raw','필수','enum','mg/dL | mmol/L','기기 설정','불허','',''),
('glucose_mgdl_std','필수(파생)','decimal(1)','mg/dL','mmol/L×18.016 (포도당 분자량 180.16 g/mol)','NA','20–600 외 플래그, CGM 범위(40–400) 상하한 표시값 별도',''),
('glucose_datetime','필수','datetime(+TZ)','ISO 8601','기기 타임스탬프','불허','미래시각·기기시계 오차 점검(동기화 로그)',''),
('glucose_type','필수','enum','공복|식후(분)|무작위|취침전|SMBG|CGM|검사실혈장','사용자 선택/기기','"미상"','미상은 민감도분석에서만',''),
('meal_relation_min','권장','int','분','식사 기록 대비 경과시간','NA','0–600',''),
('weight_kg','필수','decimal(1)','kg','체중계(블루투스) 또는 수기','NA','20–300; 1일 내 >3kg 변화 플래그','lb→kg ×0.45359237'),
('weight_datetime','필수','datetime','ISO 8601','측정 기록','불허','',''),
('weight_source','권장','enum','연동체중계|수기|검진','—','"미상"','',''),
('age_years','필수','int','세','등록 시(생년월은 키 분리보관)','불허','18–100(대상 정의에 따라)',''),
('sex','필수','enum','남|여|기타/응답거부','등록 설문','"응답거부"','',''),
('height_cm','권장','decimal','cm','측정/자가보고','NA','100–230',''),
('diabetes_type','필수','enum','없음|전당뇨|1형|2형|임신성|기타|미상','의무기록/자가보고','"미상"','',''),
('diabetes_duration_y','권장','decimal','년','진단연도','NA','0–age',''),
('treatment','필수','multi-enum','생활요법|메트포르민|SU|DPP4i|SGLT2i|GLP1RA|인슐린(기저/식전/펌프)|기타','투약 기록','"미상"','변경 시 변경일 기록',''),
('med_change_date','필수(발생 시)','date','—','투약 변경 보고/처방','NA','HbA1c 예측창과 겹침 플래그','누출·혼란 요인'),
('cgm_glucose_series','선택(확장 필수)','time series','mg/dL, 1–15분','CGM 앱/클라우드 내보내기','gap','포착률 ≥70%/14일(국제 합의) 미달 구간 표시','GMI는 별도 파생변수'),
('cgm_device_model','선택(확장 필수)','string','—','—','"미상"','센서 교체일',''),
('steps_daily','선택','int','걸음/일','스마트폰/워치','NA','착용시간 ≥10시간/일 기준 유효일',''),
('activity_minutes','선택','int','분/일','웨어러블','NA','',''),
('sleep_minutes','선택','int','분/일','웨어러블','NA','',''),
('anemia_flag / hemoglobin_gdl','권장','bool / decimal','g/dL','CBC(검사실)','NA','WHO 기준 등 연구 정의 사전명시','철결핍(HbA1c↑)·용혈(HbA1c↓) 방향 상이'),
('ckd_flag / egfr','권장','bool / decimal','mL/min/1.73m²','검사실','NA','',''),
('pregnancy_status','필수(가임 여성)','enum','비임신|임신(주수)|산후 3개월내','설문','"응답거부"','','임신 중 HbA1c 해석 상이'),
('transfusion_3mo','권장','bool','—','설문/의무기록','NA','최근 3개월 수혈 → 해당 HbA1c 분석 플래그',''),
('hemoglobinopathy','권장','enum','없음|HbS/HbC/HbE 등|미상','의무기록/검사지 경고','"미상"','방법별 간섭 확인',''),
('ethnicity','권장','enum','연구 국가 기준','설문','"응답거부"','','혈당–HbA1c 관계 집단차 평가'),
('date_shift_days','필수(공개 시)','int','일','비식별화 처리 기록(키 보관소에만)','—','개인 내 동일 shift 적용 확인','간격·순서 보존'),
]
vd=pd.DataFrame(V,columns=['변수명','수준(필수/권장/선택)','자료형','단위','출처','결측표시','품질규칙','비고'])
tim=pd.DataFrame([
('participant','1','participant_pid','인구학·당뇨유형(시간불변 or 등록시점)'),
('visit','N per participant','visit_id → participant_pid','방문일, 기관, 체중(측정형)'),
('lab_test (HbA1c)','N per participant','lab_test_id → visit_id','hba1c_date가 예측 시점(t0) 기준'),
('glucose_obs','다수(SMBG/CGM)','obs_id → device_id → participant_pid','glucose_datetime; 창 [t0−W, t0−g] 내만 입력 사용'),
('weight_obs','다수','obs_id → participant_pid','weight_datetime'),
('activity_daily','일 단위','date → participant_pid','유효착용일 플래그'),
('medication_episode','에피소드','start/end → participant_pid','예측창과 겹침 표시'),
('feature_window (파생)','HbA1c 1건당 1행','lab_test_id','입력 관측기간 W, 예측간격 g, 생성 규칙 버전'),
],columns=['테이블','카디널리티','키/연결','시간 규칙'])
with pd.ExcelWriter(str(RESEARCH_ROOT / 'report/후보_데이터셋_전체목록.xlsx'),engine='openpyxl') as w:
    df.to_excel(w,sheet_name='후보 전체',index=False)
    pd.DataFrame({'분류':['①','②','③','④'],'정의':['실제 HbA1c와 혈당·체중이 동일인에 연결','실제 HbA1c + 일부 입력(혈당 또는 체중 또는 CGM만)','HbA1c 없음 — 보조용(정답 사용 금지)','합성·출처불명·검증불가 — 정답 사용 금지']}).to_excel(w,sheet_name='분류 정의',index=False)
    pd.DataFrame({'확인수준':['원자료확인','공식변수사전확인','문헌확인','미확인'],'의미':['이번에 원자료를 받아 직접 집계','공식 코드북/사전에서 변수 확인(원자료 미확보)','원논문·공식 안내문 서술로 확인','이번 조사에서 근거 미확보']}).to_excel(w,sheet_name='확인수준 정의',index=False)
with pd.ExcelWriter(str(RESEARCH_ROOT / 'report/변수사전.xlsx'),engine='openpyxl') as w:
    vd.to_excel(w,sheet_name='변수사전',index=False)
    tim.to_excel(w,sheet_name='시간연결구조',index=False)
    pd.DataFrame([('HbA1c %→mmol/mol','IFCC = 10.93 × NGSP − 23.50','NGSP/IFCC 마스터방정식 (ngsp.org/ifcc.asp)'),('HbA1c mmol/mol→%','NGSP = 0.09148 × IFCC + 2.152','동일'),('혈당 mmol/L→mg/dL','mg/dL = mmol/L × 18.016','포도당 분자량 180.16 g/mol; NHANES는 mg/dL×0.05551'),('GMI(파생, 정답 아님)','GMI% = 3.31 + 0.02392 × 평균 CGM 혈당(mg/dL)','Bergenstal et al. Diabetes Care 2018'),('eAG(파생, 정답 아님)','eAG(mg/dL) = 28.7 × HbA1c% − 46.7','ADAG (Nathan 2008)'),('체중 lb→kg','kg = lb × 0.45359237','국제 정의')],columns=['변환','공식','근거']).to_excel(w,sheet_name='단위변환',index=False)
for p in [str(RESEARCH_ROOT / 'report/후보_데이터셋_전체목록.xlsx'),str(RESEARCH_ROOT / 'report/변수사전.xlsx')]:
    from openpyxl import load_workbook
    wb=load_workbook(p)
    for ws in wb.worksheets:
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='305496')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or ''))*0.9 for c in col[:30])))
            for c in col: c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.freeze_panes='C2' if ws.title=='후보 전체' else 'A2'
    wb.save(p)
df.to_csv(str(RESEARCH_ROOT / 'report/후보_데이터셋_전체목록.csv'),index=False,encoding='utf-8-sig')
print(len(df),'candidates;',len(vd),'variables')
